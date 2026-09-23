from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import Comprobante, MovimientoCaja, MovimientoCombustible
from app.services.comprobante_calc import ESTADO_LABELS, TIPO_LABELS

_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="0F3D2E")
_TOTAL_FONT = Font(name="Calibri", bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="ECFDF5")
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)
_THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
_MONEY = '#,##0.00'
_GAL = '#,##0.000'


def _nueva_hoja(titulo: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 22
    return wb, ws


def _anchos(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _guardar(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _fecha(value) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value or "")


def _items_texto(comprobante: Comprobante) -> str:
    partes = []
    for item in getattr(comprobante, "items", None) or []:
        desc = (getattr(item, "descripcion", None) or "").strip()
        if not desc:
            continue
        precio = float(getattr(item, "precio_unitario", 0) or 0)
        cant = float(getattr(item, "cantidad", 1) or 1)
        qty = ""
        if abs(cant - 1) > 0.0001:
            qty_txt = f"{cant:,.3f}".rstrip("0").rstrip(".")
            qty = f" x {qty_txt}"
        partes.append(f"{desc} - S/ {precio:,.2f}{qty}")
    return "\n".join(partes)


def generar_excel_reporte_comprobantes(docs: list[Comprobante]) -> bytes:
    headers = [
        "Fecha",
        "Tipo",
        "Serie-Nº",
        "Cliente",
        "Documento",
        "Zona",
        "Motivo",
        "Descripción",
        "Estado",
        "Total",
    ]
    wb, ws = _nueva_hoja("Comprobantes", headers)
    total_general = 0.0
    for doc in docs:
        total = float(doc.total or 0)
        total_general += total
        descripcion = _items_texto(doc)
        ws.append(
            [
                _fecha(doc.fecha_emision),
                TIPO_LABELS.get(doc.tipo, str(doc.tipo)),
                f"{doc.serie}-{doc.numero}",
                doc.cliente_nombre or "",
                doc.cliente_documento or "",
                getattr(doc, "zona", None) or "",
                getattr(doc, "motivo", None) or "",
                descripcion,
                ESTADO_LABELS.get(doc.estado.value, doc.estado.value),
                total,
            ]
        )
        row = ws.max_row
        lineas = max(descripcion.count("\n") + 1, 1) if descripcion else 1
        ws.row_dimensions[row].height = max(18, 15 * lineas)
        for cell in ws[row]:
            cell.alignment = _WRAP
            cell.border = _THIN
        ws.cell(row, 10).number_format = _MONEY

    total_row = ws.max_row + 1
    ws.cell(total_row, 9, "TOTAL").font = _TOTAL_FONT
    total_cell = ws.cell(total_row, 10, total_general)
    total_cell.font = _TOTAL_FONT
    total_cell.number_format = _MONEY
    for cell in ws[total_row]:
        cell.fill = _TOTAL_FILL
        cell.border = _THIN
    _anchos(ws, [14, 24, 16, 32, 16, 18, 28, 42, 14, 14])
    return _guardar(wb)


def generar_excel_reporte_cajas(movimientos: list[MovimientoCaja]) -> bytes:
    headers = ["Fecha", "Caja", "Tipo", "N° Transacción", "Concepto", "Monto"]
    wb, ws = _nueva_hoja("Cajas", headers)
    ingresos = 0.0
    egresos = 0.0
    for mov in movimientos:
        tipo_val = str(getattr(mov.tipo, "value", mov.tipo))
        monto = float(mov.monto or 0)
        if tipo_val == "egreso":
            egresos += monto
            monto_celda = -monto
            tipo_lbl = "Egreso"
        else:
            ingresos += monto
            monto_celda = monto
            tipo_lbl = "Ingreso"
        caja_nombre = mov.caja.nombre if getattr(mov, "caja", None) else ""
        ws.append(
            [
                _fecha(mov.fecha),
                caja_nombre,
                tipo_lbl,
                getattr(mov, "numero_transaccion", None) or "",
                mov.concepto or "",
                monto_celda,
            ]
        )
        row = ws.max_row
        for cell in ws[row]:
            cell.alignment = _WRAP
            cell.border = _THIN
        ws.cell(row, 6).number_format = _MONEY

    total_row = ws.max_row + 1
    ws.cell(total_row, 4, "Ingresos").font = _TOTAL_FONT
    ing_cell = ws.cell(total_row, 5, ingresos)
    ing_cell.number_format = _MONEY
    ing_cell.font = _TOTAL_FONT
    ws.cell(total_row + 1, 4, "Egresos").font = _TOTAL_FONT
    egr_cell = ws.cell(total_row + 1, 5, egresos)
    egr_cell.number_format = _MONEY
    egr_cell.font = _TOTAL_FONT
    ws.cell(total_row + 2, 4, "Saldo").font = _TOTAL_FONT
    saldo_cell = ws.cell(total_row + 2, 5, ingresos - egresos)
    saldo_cell.number_format = _MONEY
    saldo_cell.font = _TOTAL_FONT
    for row in range(total_row, total_row + 3):
        for cell in ws[row]:
            cell.fill = _TOTAL_FILL
            cell.border = _THIN
    _anchos(ws, [14, 24, 12, 22, 40, 14])
    return _guardar(wb)


def generar_excel_reporte_combustibles(movimientos: list[MovimientoCombustible]) -> bytes:
    headers = ["Fecha", "Tipo", "Galones", "Conductor", "Marca", "Placa", "Notas"]
    wb, ws = _nueva_hoja("Combustibles", headers)
    ingresos = 0.0
    salidas = 0.0
    for mov in movimientos:
        tipo_val = str(getattr(mov.tipo, "value", mov.tipo))
        galones = float(mov.galones or 0)
        if tipo_val == "salida":
            salidas += galones
            galones_celda = -galones
            tipo_lbl = "Salida"
        else:
            ingresos += galones
            galones_celda = galones
            tipo_lbl = "Ingreso"
        ws.append(
            [
                _fecha(mov.fecha),
                tipo_lbl,
                galones_celda,
                mov.conductor or "",
                mov.marca or "",
                mov.placa or "",
                mov.notas or "",
            ]
        )
        row = ws.max_row
        for cell in ws[row]:
            cell.alignment = _WRAP
            cell.border = _THIN
        ws.cell(row, 3).number_format = _GAL

    total_row = ws.max_row + 1
    ws.cell(total_row, 1, "Ingresos").font = _TOTAL_FONT
    ing_cell = ws.cell(total_row, 3, ingresos)
    ing_cell.number_format = _GAL
    ing_cell.font = _TOTAL_FONT
    ws.cell(total_row + 1, 1, "Salidas").font = _TOTAL_FONT
    sal_cell = ws.cell(total_row + 1, 3, salidas)
    sal_cell.number_format = _GAL
    sal_cell.font = _TOTAL_FONT
    ws.cell(total_row + 2, 1, "Saldo").font = _TOTAL_FONT
    saldo_cell = ws.cell(total_row + 2, 3, ingresos - salidas)
    saldo_cell.number_format = _GAL
    saldo_cell.font = _TOTAL_FONT
    for row in range(total_row, total_row + 3):
        for cell in ws[row]:
            cell.fill = _TOTAL_FILL
            cell.border = _THIN
    _anchos(ws, [14, 12, 14, 28, 18, 14, 36])
    return _guardar(wb)
